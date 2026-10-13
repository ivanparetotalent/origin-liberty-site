"""Minimal markdown -> .docx for the checklist content file."""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor

src, out = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)

def runs(p, text):
    for i, part in enumerate(re.split(r'\*\*', text)):
        if part:
            r = p.add_run(part); r.bold = (i % 2 == 1)

lines = open(src, encoding='utf-8').read().splitlines()
i = 0
while i < len(lines):
    ln = lines[i]
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r'-+', c) for c in cells): rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c); cell.text = ''
                runs(cell.paragraphs[0], f'**{val}**' if r == 0 else val)
        continue
    if ln.strip() == '---':
        doc.add_paragraph('')
    elif ln.startswith('### '): doc.add_heading(ln[4:], level=3)
    elif ln.startswith('## '): doc.add_heading(ln[3:], level=2)
    elif ln.startswith('# '): doc.add_heading(ln[2:], level=1)
    elif re.match(r'^  +- ', ln): runs(doc.add_paragraph(style='List Bullet 2'), ln.strip()[2:])
    elif ln.startswith('- '): runs(doc.add_paragraph(style='List Bullet'), ln[2:])
    elif re.match(r'^\d+\. ', ln): runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', ln))
    elif ln.strip(): runs(doc.add_paragraph(), ln)
    i += 1
doc.save(out); print('saved', out)
